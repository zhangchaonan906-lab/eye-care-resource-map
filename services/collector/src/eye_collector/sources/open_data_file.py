from __future__ import annotations

import csv
import hashlib
import io
import re
import stat
import warnings
import zipfile
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any

from eye_collector.hashing import canonical_sha256
from eye_collector.models import RawRecord, SourceDescriptor, SourcePage
from eye_collector.sources.base import SourceAdapter

_MAX_FILE_BYTES = 10 * 1024 * 1024
_MAX_ARCHIVE_MEMBER_BYTES = 10 * 1024 * 1024
_MAX_RECORDS_PER_RUN = 150


@dataclass(frozen=True, slots=True)
class OpenDataFileInspection:
    original_filename: str
    file_sha256: str
    file_size_bytes: int
    detected_format: str
    headers: tuple[str, ...]
    row_count: int
    expected_schema: tuple[str, ...]
    schema_match: bool
    container_format: str | None = None
    archive_member_name: str | None = None
    archive_member_sha256: str | None = None
    archive_member_size_bytes: int | None = None
    validation_errors: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class _ArchiveMember:
    name: str
    content: bytes


def _zip_xlsx_member(path: Path, dataset: OpenDataDataset) -> _ArchiveMember:
    if not dataset.allow_zip_xlsx:
        raise ValueError("ZIP input is not approved for this dataset")
    try:
        with zipfile.ZipFile(path, "r") as archive:
            members = archive.infolist()
            if len(members) != 1:
                raise ValueError("ZIP must contain exactly one data file")
            member = members[0]
            if member.flag_bits & 0x1:
                raise ValueError("encrypted ZIP members are not supported")
            member_path = PurePosixPath(member.filename)
            mode = member.external_attr >> 16
            if (
                not member.filename
                or member.filename.startswith(("/", "\\"))
                or "\\" in member.filename
                or ":" in member.filename.split("/", 1)[0]
                or "\x00" in member.filename
                or member_path.is_absolute()
                or any(part in {".", ".."} for part in member_path.parts)
                or member.is_dir()
                or stat.S_ISLNK(mode)
                or mode & 0o111
            ):
                raise ValueError("ZIP contains an unsafe path or executable member")
            if member_path.suffix.lower() != ".xlsx":
                raise ValueError("ZIP must contain one XLSX data file")
            if member.file_size < 1 or member.file_size > _MAX_ARCHIVE_MEMBER_BYTES:
                raise ValueError("ZIP XLSX member must be no larger than 10 MiB")
            with archive.open(member, "r") as member_stream:
                content = member_stream.read(_MAX_ARCHIVE_MEMBER_BYTES + 1)
            if len(content) > _MAX_ARCHIVE_MEMBER_BYTES:
                raise ValueError("ZIP XLSX member must be no larger than 10 MiB")
            if len(content) != member.file_size:
                raise ValueError("ZIP member size does not match its directory entry")
            return _ArchiveMember(member.filename, content)
    except (zipfile.BadZipFile, OSError, RuntimeError) as error:
        raise ValueError("invalid or unreadable ZIP archive") from error


def _mapped_values_valid(
    headers: Sequence[str], row: Sequence[object], dataset: OpenDataDataset, row_number: int
) -> tuple[str, ...]:
    administrative_header = next(
        (header for header, target in dataset.field_mapping if target == "administrative_code"),
        None,
    )
    if administrative_header is None or administrative_header not in headers:
        return ()
    value = row[headers.index(administrative_header)]
    text = _string_value(value)
    if text is not None and re.fullmatch(r"[0-9]{6}", text) is None:
        return (
            f"row {row_number}: {administrative_header} must be a six-digit administrative code",
        )
    return ()


def inspect_open_data_file(path: str | Path, dataset: OpenDataDataset) -> OpenDataFileInspection:
    """Read-only file preflight; hashes raw bytes and checks the full tabular schema."""
    source_path = Path(path)
    if not source_path.is_file():
        raise ValueError("source file does not exist")
    before = source_path.stat()
    if before.st_size > _MAX_FILE_BYTES:
        raise ValueError("source file size exceeds the 10 MiB limit")
    if before.st_size == 0:
        raise ValueError("source file is empty")
    original_sha256 = sha256_file(source_path)

    suffix = source_path.suffix.lower()
    with source_path.open("rb") as stream:
        signature = stream.read(8)
    archive_member: _ArchiveMember | None = None
    container_format: str | None = None
    if suffix == ".zip":
        archive_member = _zip_xlsx_member(source_path, dataset)
        detected_format = "zip"
        container_format = "zip"
    elif suffix == ".xlsx" and signature.startswith(b"PK\x03\x04"):
        detected_format = "xlsx"
    elif suffix == ".xls" and signature.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        detected_format = "xls"
    elif suffix == ".csv" and not signature.startswith(
        (b"PK\x03\x04", bytes.fromhex("D0CF11E0A1B11AE1"))
    ):
        detected_format = "csv"
    else:
        raise ValueError("file extension and detected format do not match")

    inspection_source: str | Path | io.BytesIO = (
        io.BytesIO(archive_member.content) if archive_member else source_path
    )
    if detected_format == "csv":
        headers, row_count, widths_match, validation_errors = _inspect_csv(source_path, dataset)
    elif detected_format in {"xlsx", "zip"}:
        headers, row_count, widths_match, validation_errors = _inspect_xlsx(
            inspection_source, dataset
        )
    else:
        headers, row_count, widths_match, validation_errors = _inspect_xls(source_path, dataset)

    after = source_path.stat()
    if (
        before.st_size != after.st_size
        or before.st_mtime_ns != after.st_mtime_ns
        or original_sha256 != sha256_file(source_path)
    ):
        raise ValueError("source file changed while it was being inspected")
    header_set_matches = len(headers) == len(set(headers)) and set(headers) == set(
        dataset.expected_headers
    )
    return OpenDataFileInspection(
        original_filename=source_path.name,
        file_sha256=original_sha256,
        file_size_bytes=before.st_size,
        detected_format=detected_format,
        headers=headers,
        row_count=row_count,
        expected_schema=dataset.expected_headers,
        schema_match=header_set_matches and widths_match and not validation_errors,
        container_format=container_format,
        archive_member_name=archive_member.name if archive_member else None,
        archive_member_sha256=(
            hashlib.sha256(archive_member.content).hexdigest() if archive_member else None
        ),
        archive_member_size_bytes=len(archive_member.content) if archive_member else None,
        validation_errors=validation_errors,
    )


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(64 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _inspect_csv(
    path: Path, dataset: OpenDataDataset
) -> tuple[tuple[str, ...], int, bool, tuple[str, ...]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        for _ in range(dataset.header_row - 1):
            next(reader, None)
        raw_headers = next(reader, None)
        if raw_headers is None:
            raise ValueError("source file header schema is missing")
        headers = tuple(raw_headers)
        row_count = 0
        widths_match = True
        validation_errors: list[str] = []
        for row in reader:
            if not row or all(value == "" for value in row):
                continue
            row_count += 1
            if len(row) != len(headers):
                widths_match = False
            else:
                validation_errors.extend(_mapped_values_valid(headers, row, dataset, row_count + 1))
    return headers, row_count, widths_match, tuple(validation_errors)


def _inspect_xlsx(
    source: str | Path | io.BytesIO, dataset: OpenDataDataset
) -> tuple[tuple[str, ...], int, bool, tuple[str, ...]]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
        )
        workbook = load_workbook(source, read_only=False, data_only=True)
    try:
        if dataset.data_sheet_name is not None:
            if dataset.data_sheet_name not in workbook.sheetnames:
                raise ValueError(f"configured worksheet {dataset.data_sheet_name!r} is missing")
            sheet = workbook[dataset.data_sheet_name]
        elif len(workbook.worksheets) == 1:
            sheet = workbook.worksheets[0]
        else:
            raise ValueError("workbook has multiple sheets; a data_sheet_name is required")
        row_iter = sheet.iter_rows(values_only=True)
        for _ in range(dataset.header_row - 1):
            next(row_iter, None)
        raw_headers = next(row_iter, None)
        if raw_headers is None:
            raise ValueError("source file header schema is missing")
        headers = tuple("" if value is None else str(value) for value in raw_headers)
        row_count = 0
        widths_match = True
        validation_errors: list[str] = []
        for row in row_iter:
            if all(value is None or value == "" for value in row):
                continue
            row_count += 1
            if len(row) != len(headers):
                widths_match = False
            else:
                validation_errors.extend(_mapped_values_valid(headers, row, dataset, row_count + 1))
        return headers, row_count, widths_match, tuple(validation_errors)
    finally:
        workbook.close()


def _inspect_xls(
    path: Path, dataset: OpenDataDataset
) -> tuple[tuple[str, ...], int, bool, tuple[str, ...]]:
    import xlrd  # type: ignore[import-untyped]

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        if dataset.data_sheet_name is not None:
            try:
                sheet = workbook.sheet_by_name(dataset.data_sheet_name)
            except xlrd.biffh.XLRDError as error:
                raise ValueError(
                    f"configured worksheet {dataset.data_sheet_name!r} is missing"
                ) from error
        elif workbook.nsheets == 1:
            sheet = workbook.sheet_by_index(0)
        else:
            raise ValueError("workbook has multiple sheets; a data_sheet_name is required")
        if sheet.nrows == 0:
            raise ValueError("source file header schema is missing")
        header_index = dataset.header_row - 1
        if header_index >= sheet.nrows:
            raise ValueError("source file header schema is missing")
        headers = tuple(str(value) for value in sheet.row_values(header_index))
        row_count = 0
        validation_errors: list[str] = []
        for row_index in range(header_index + 1, sheet.nrows):
            values = sheet.row_values(row_index)
            if any(value is not None and value != "" for value in values):
                row_count += 1
                validation_errors.extend(
                    _mapped_values_valid(headers, values, dataset, row_index + 1)
                )
        return headers, row_count, True, tuple(validation_errors)
    finally:
        workbook.release_resources()


@dataclass(frozen=True, slots=True)
class OpenDataDataset:
    source_key: str
    source_name: str
    dataset_url: str
    expected_headers: tuple[str, ...]
    field_mapping: tuple[tuple[str, str], ...]
    stable_key_header: str | None = None
    data_sheet_name: str | None = None
    allow_zip_xlsx: bool = False
    header_row: int = 1

    def __post_init__(self) -> None:
        if not isinstance(self.header_row, int) or isinstance(self.header_row, bool):
            raise ValueError("header_row must be a positive one-based row number")
        if self.header_row < 1:
            raise ValueError("header_row must be a positive one-based row number")
        raw_headers = [header for header, _canonical in self.field_mapping]
        if len(raw_headers) != len(set(raw_headers)):
            raise ValueError("field mapping contains duplicate source headers")
        if not set(raw_headers).issubset(self.expected_headers):
            raise ValueError("mapped source headers must be in expected_headers")
        if (
            self.stable_key_header is not None
            and self.stable_key_header not in self.expected_headers
        ):
            raise ValueError("stable_key_header must be in expected_headers")
        if not self.dataset_url.startswith("https://"):
            raise ValueError("dataset_url must use HTTPS")


BEIJING_HOSPITALS = OpenDataDataset(
    source_key="beijing-open-data-hospitals",
    source_name="北京市公共数据开放平台-医院",
    dataset_url="https://data.beijing.gov.cn/zyml/wnkfsj/5652.htm",
    expected_headers=("机构名称",),
    field_mapping=(("机构名称", "name"),),
    stable_key_header="机构名称",
)

TIANJIN_REGISTRATION_PREVIEW = OpenDataDataset(
    source_key="tianjin-registration-local-preview",
    source_name="天津市卫生健康委-医疗机构执业登记信息（本地预演）",
    dataset_url="https://open.data.tj.gov.cn/sjj/8e3f7e670ea9492dbc480e2c68683ce5.htm",
    expected_headers=("批准时间", "机构名称", "地址", "诊疗科目", "床位数", "类别", "所有制形式"),
    field_mapping=(
        ("批准时间", "approval_date"),
        ("机构名称", "name"),
        ("地址", "address"),
        ("诊疗科目", "specialties"),
        ("床位数", "bed_count"),
        ("类别", "source_category"),
        ("所有制形式", "ownership_type"),
    ),
    data_sheet_name="医疗机构执业登记",
    header_row=2,
)

BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS = OpenDataDataset(
    source_key="beijing-open-data-designated-medical-institutions",
    source_name="北京市公共数据开放平台-定点医疗机构信息",
    dataset_url="https://data.beijing.gov.cn/zyml/ajg/sybj/17425.htm",
    expected_headers=(
        "序号",
        "医院名称",
        "定点医疗机构编码",
        "所属区",
        "医院类别",
        "医院等级",
        "医院地址",
        "数据唯一记录号",
        "数据创建时间",
        "数据更新时间",
    ),
    field_mapping=(
        ("医院名称", "name"),
        ("医院地址", "address"),
        ("所属区", "administrative_code"),
        ("医院等级", "hospital_grade"),
        ("医院类别", "source_category"),
        ("定点医疗机构编码", "registration_id"),
    ),
    stable_key_header="定点医疗机构编码",
    data_sheet_name="北京市医疗保障局-定点医疗机构信息",
)

SHENZHEN_BAOAN_HOSPITALS = OpenDataDataset(
    source_key="shenzhen-open-data-baoan-hospital-basic-information",
    source_name="深圳市政府数据开放平台-宝安区-医院基本信息",
    dataset_url="https://opendata.sz.gov.cn/data/dataSet/toDataDetails/29200_02800636",
    expected_headers=(
        "文档ID",
        "名称",
        "所在区县",
        "主管部门",
        "详细地址",
        "邮政编码",
        "联系电话",
        "电子邮箱",
        "网站地址",
        "性质",
        "级别",
        "等级",
        "是否医保指定医院",
        "医院简介",
        "所获表彰与荣誉",
        "医疗优势与特长",
        "交通情况",
    ),
    field_mapping=(
        ("文档ID", "source_reference_id"),
        ("名称", "name"),
        ("所在区县", "administrative_context"),
        ("详细地址", "address"),
        ("性质", "source_category"),
        ("级别", "hospital_level"),
        ("等级", "hospital_grade"),
        ("医疗优势与特长", "specialties"),
        ("医院简介", "hospital_description"),
    ),
    stable_key_header="文档ID",
    data_sheet_name="数据集1",
    allow_zip_xlsx=True,
)


class OpenDataFileAdapter(SourceAdapter):
    """Read an operator-downloaded official CSV/XLS/XLSX without making network calls."""

    def __init__(self, path: str | Path, dataset: OpenDataDataset) -> None:
        self._path = Path(path)
        self._dataset = dataset
        supported = {".csv", ".xls", ".xlsx"}
        if dataset.allow_zip_xlsx:
            supported.add(".zip")
        if self._path.suffix.lower() not in supported:
            raise ValueError("supported formats are CSV, XLS, XLSX, and approved ZIP")
        if not self._path.is_file():
            raise ValueError("source file does not exist")
        if self._path.stat().st_size > _MAX_FILE_BYTES:
            raise ValueError("source file size exceeds the 10 MiB limit")
        self.original_filename = self._path.name
        self.file_size_bytes = self._path.stat().st_size
        self.file_sha256 = sha256_file(self._path)
        self.archive_member_name: str | None = None
        self.archive_member_sha256: str | None = None
        self.archive_member_size_bytes: int | None = None
        self._archive_member_content: bytes | None = None
        if self._path.suffix.lower() == ".zip":
            member = _zip_xlsx_member(self._path, dataset)
            self.archive_member_name = member.name
            self.archive_member_sha256 = hashlib.sha256(member.content).hexdigest()
            self.archive_member_size_bytes = len(member.content)
            self._archive_member_content = member.content

    @property
    def descriptor(self) -> SourceDescriptor:
        return SourceDescriptor(
            source_key=self._dataset.source_key,
            source_name=self._dataset.source_name,
            catalog_url=self._dataset.dataset_url,
            access_method="file",
        )

    def fetch_page(self, region_code: str, cursor: str | None) -> SourcePage:
        if cursor is not None:
            raise ValueError("file adapter does not support pagination")
        if len(region_code) != 6 or not region_code.isdigit():
            raise ValueError("region_code must contain six digits")
        return SourcePage(self._read_records(_MAX_RECORDS_PER_RUN), None)

    def preview_records(self) -> tuple[RawRecord, ...]:
        """Read local rows for an in-memory QA preview, including rows without a name."""
        return self._read_records(_MAX_RECORDS_PER_RUN, require_name=False)

    def iter_pages(
        self,
        region_code: str,
        limit: int | None = None,
        *,
        on_request: Callable[[], None] | None = None,
    ) -> Iterator[SourcePage]:
        if limit is None:
            raise ValueError("official file imports require an explicit record limit")
        if limit > _MAX_RECORDS_PER_RUN:
            raise ValueError("official file imports are limited to 150 records per run")
        if len(region_code) != 6 or not region_code.isdigit():
            raise ValueError("region_code must contain six digits")
        if limit < 1:
            raise ValueError("limit must be a positive integer")
        if on_request is not None:
            on_request()
        yield SourcePage(self._read_records(limit), None)

    def _read_records(self, limit: int, *, require_name: bool = True) -> tuple[RawRecord, ...]:
        self._assert_fingerprint_unchanged()
        suffix = self._path.suffix.lower()
        if suffix == ".csv":
            rows = self._read_csv(limit)
        elif suffix in {".xlsx", ".zip"}:
            rows = self._read_xlsx(limit)
        else:
            rows = self._read_xls(limit)

        records: list[RawRecord] = []
        for row in rows:
            if all(value is None or value == "" for value in row.values()):
                continue
            source_fields = {
                header: _string_value(row[header]) for header, _name in self._dataset.field_mapping
            }
            for source_header, canonical_name in self._dataset.field_mapping:
                if canonical_name == "administrative_code":
                    district_code = source_fields[source_header]
                    if (
                        district_code is not None
                        and re.fullmatch(r"[0-9]{6}", district_code) is None
                    ):
                        raise ValueError(f"{source_header} must be a six-digit administrative code")
            mapped = {
                canonical_name: source_fields[source_header]
                for source_header, canonical_name in self._dataset.field_mapping
                if source_fields[source_header] is not None
            }
            name = mapped.get("name")
            if require_name and (not isinstance(name, str) or not name.strip()):
                raise ValueError("source row is missing a required hospital name")
            stable_value = row.get(self._dataset.stable_key_header or "")
            source_key = _string_value(stable_value)
            if not source_key:
                source_key = canonical_sha256(source_fields)
            payload: dict[str, Any] = {"source_fields": source_fields, **mapped}
            records.append(RawRecord(source_key, self._dataset.dataset_url, payload))
            if len(records) >= limit:
                break
        self._assert_fingerprint_unchanged()
        return tuple(records)

    def _assert_fingerprint_unchanged(self) -> None:
        if (
            self._path.stat().st_size != self.file_size_bytes
            or sha256_file(self._path) != self.file_sha256
        ):
            raise ValueError("source file changed after fingerprint capture")

    def _read_csv(self, limit: int) -> list[dict[str, object]]:
        with self._path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            for _ in range(self._dataset.header_row - 1):
                next(reader, None)
            headers = next(reader, None)
            if headers is None:
                raise ValueError("source file header schema is missing")
            row_reader = csv.DictReader(stream, fieldnames=headers)
            self._validate_headers(headers)
            rows: list[dict[str, object]] = []
            for row in row_reader:
                if None in row or any(value is None for value in row.values()):
                    raise ValueError("source file row width does not match its approved headers")
                if all(value is None or value == "" for value in row.values()):
                    continue
                rows.append(dict(row))
                if len(rows) >= limit:
                    break
            return rows

    def _read_xlsx(self, limit: int) -> list[dict[str, object]]:
        from openpyxl import load_workbook

        source: str | Path | io.BytesIO = (
            io.BytesIO(self._archive_member_content)
            if self._archive_member_content is not None
            else self._path
        )
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style, apply openpyxl's default",
                category=UserWarning,
            )
            workbook = load_workbook(source, read_only=False, data_only=True)
        try:
            if self._dataset.data_sheet_name is not None:
                if self._dataset.data_sheet_name not in workbook.sheetnames:
                    raise ValueError(
                        f"configured worksheet {self._dataset.data_sheet_name!r} is missing"
                    )
                sheet = workbook[self._dataset.data_sheet_name]
            elif len(workbook.worksheets) == 1:
                sheet = workbook.worksheets[0]
            else:
                raise ValueError("workbook has multiple sheets; a data_sheet_name is required")
            row_iter = sheet.iter_rows(values_only=True)
            for _ in range(self._dataset.header_row - 1):
                next(row_iter, None)
            headers = next(row_iter, None)
            names = self._validate_headers(headers)
            rows: list[dict[str, object]] = []
            for values in row_iter:
                if all(value is None or value == "" for value in values):
                    continue
                rows.append(dict(zip(names, values, strict=True)))
                if len(rows) >= limit:
                    break
            return rows
        finally:
            workbook.close()

    def _read_xls(self, limit: int) -> list[dict[str, object]]:
        import xlrd

        workbook = xlrd.open_workbook(self._path, on_demand=True)
        try:
            if self._dataset.data_sheet_name is not None:
                try:
                    sheet = workbook.sheet_by_name(self._dataset.data_sheet_name)
                except xlrd.biffh.XLRDError as error:
                    raise ValueError(
                        f"configured worksheet {self._dataset.data_sheet_name!r} is missing"
                    ) from error
            elif workbook.nsheets == 1:
                sheet = workbook.sheet_by_index(0)
            else:
                raise ValueError("workbook has multiple sheets; a data_sheet_name is required")
            if sheet.nrows == 0:
                raise ValueError("source file header schema is missing")
            header_index = self._dataset.header_row - 1
            if header_index >= sheet.nrows:
                raise ValueError("source file header schema is missing")
            headers = self._validate_headers(sheet.row_values(header_index))
            rows: list[dict[str, object]] = []
            for row_index in range(header_index + 1, sheet.nrows):
                row: dict[str, object] = {}
                for column_index, header in enumerate(headers):
                    cell = sheet.cell(row_index, column_index)
                    if (
                        header == self._dataset.stable_key_header
                        and cell.ctype == xlrd.XL_CELL_NUMBER
                    ):
                        workbook.release_resources()
                        raise ValueError("numeric stable IDs in XLS are ambiguous; use CSV/XLSX")
                    row[header] = cell.value
                if all(value is None or value == "" for value in row.values()):
                    continue
                rows.append(row)
                if len(rows) >= limit:
                    break
            return rows
        finally:
            workbook.release_resources()

    def _validate_headers(self, headers: Sequence[object] | None) -> list[str]:
        if headers is None or any(header is None for header in headers):
            raise ValueError("source file header schema is missing")
        actual = [str(header) for header in headers]
        if len(actual) != len(set(actual)) or set(actual) != set(self._dataset.expected_headers):
            raise ValueError(
                "source file header schema does not match the approved dataset mapping"
            )
        return actual


def _string_value(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value)
    return text if text else None
