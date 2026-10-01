from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from eye_collector.tianjin_preview import preview_tianjin_registration_file


def test_tianjin_preview_reports_quality_and_p3_evidence_without_persistence(
    tmp_path: Path,
) -> None:
    source_file = tmp_path / "synthetic-tianjin-registration.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "医疗机构执业登记"
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "天津医疗机构执业登记信息"
    sheet.append(
        ["批准时间", "机构名称", "地址", "诊疗科目", "床位数", "类别", "所有制形式"]
    )
    sheet.append(
        ["2026-01-02", "天津测试医院A", "天津市西青区测试路1号", "眼科", 20, "综合医院", "公立"]
    )
    sheet.append(
        [
            "2026-01-03",
            "天津测试眼科医院B",
            "天津市南开区测试路2号",
            "内科",
            "not-a-number",
            "三级眼科医院",
            "公立",
        ]
    )
    sheet.append(["2026-01-04", "", "天津市河西区测试路3号", "眼科", None, "护理院", "私人"])
    sheet.append(
        [
            "not-a-date",
            "天津测试医院A",
            "天津市西青区测试路1号；天津市南开区测试路4号",
            "内科",
            15,
            "综合医院",
            "公立",
        ]
    )
    workbook.create_sheet("Sheet3")
    workbook.save(source_file)

    report = preview_tianjin_registration_file(source_file)

    assert report.total_raw_rows == 4
    assert report.valid_names == 3
    assert report.empty_names == 1
    assert report.empty_addresses == 0
    assert report.duplicate_exact_name_groups == 1
    assert report.duplicate_normalized_name_groups == 1
    assert report.distinct_normalized_names == 2
    assert report.candidate_groups == 2
    assert report.rows_with_eye_evidence == 1
    assert report.distinct_eye_candidate_names == 1
    assert dict(report.category_distribution) == {
        "三级眼科医院": 1,
        "护理院": 1,
        "综合医院": 2,
    }
    assert dict(report.ownership_distribution) == {"公立": 3, "私人": 1}
    assert (report.beds_valid, report.beds_missing, report.beds_invalid) == (2, 1, 1)
    assert report.approval_date_min == "2026-01-02"
    assert report.approval_date_max == "2026-01-04"
    assert report.approval_date_invalid == 1
    assert report.observed_districts == ("南开区", "河西区", "西青区")
    assert report.multi_location_address_candidates == 1
