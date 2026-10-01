from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path

import pytest
import xlwt
from openpyxl import Workbook

from eye_collector.sources.open_data_file import (
    BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS,
    BEIJING_HOSPITALS,
    SHENZHEN_BAOAN_HOSPITALS,
    OpenDataFileInspection,
    inspect_open_data_file,
)

BEIJING_HEADERS = (
    "序号",
    "医院名称",
    "定点医疗机构编码",
    "所属区",
    "医院类别",
    "医院等级",
    "医院地址",
    "数据唯一记录号",
    "数据创建时间",
    "数据更新时间",
)
SHENZHEN_HEADERS = (
    "文档ID",
    "名称",
    "所在区县",
    "主管部门",
    "详细地址",
    "邮政编码",
    "联系电话",
    "电子邮箱",
    "网站地址",
    "性质",
    "级别",
    "等级",
    "是否医保指定医院",
    "医院简介",
    "所获表彰与荣誉",
    "医疗优势与特长",
    "交通情况",
)


def xlsx_bytes(sheet_name: str, headers: tuple[str, ...], rows: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def write_zip(path: Path, members: dict[str, bytes], *, encrypted_flag: bool = False) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            info = zipfile.ZipInfo(name)
            if encrypted_flag:
                info.flag_bits |= 0x1
            archive.writestr(info, content)
    if encrypted_flag:
        content = bytearray(path.read_bytes())
        local_offset = content.find(b"PK\x03\x04")
        central_offset = content.rfind(b"PK\x01\x02")
        for flag_offset in (local_offset + 6, central_offset + 8):
            flags = int.from_bytes(content[flag_offset : flag_offset + 2], "little")
            content[flag_offset : flag_offset + 2] = (flags | 0x1).to_bytes(2, "little")
        path.write_bytes(content)


def test_preflight_hashes_original_bytes_and_counts_all_csv_rows(tmp_path: Path) -> None:
    source_file = tmp_path / "official.csv"
    original_bytes = "机构名称\r\n医院甲\r\n医院乙\r\n".encode("utf-8-sig")
    source_file.write_bytes(original_bytes)

    inspection = inspect_open_data_file(source_file, BEIJING_HOSPITALS)

    assert inspection == OpenDataFileInspection(
        original_filename="official.csv",
        file_sha256=hashlib.sha256(original_bytes).hexdigest(),
        file_size_bytes=len(original_bytes),
        detected_format="csv",
        headers=("机构名称",),
        row_count=2,
        expected_schema=("机构名称",),
        schema_match=True,
    )
    assert source_file.read_bytes() == original_bytes


def test_preflight_reports_schema_mismatch_without_changing_file(tmp_path: Path) -> None:
    source_file = tmp_path / "changed.csv"
    source_file.write_text("机构名称,意外字段\n医院甲,值\n", encoding="utf-8-sig")
    before = source_file.read_bytes()

    inspection = inspect_open_data_file(source_file, BEIJING_HOSPITALS)

    assert inspection.schema_match is False
    assert inspection.headers == ("机构名称", "意外字段")
    assert inspection.row_count == 1
    assert source_file.read_bytes() == before


@pytest.mark.parametrize("extension", [".xls", ".xlsx"])
def test_preflight_detects_excel_format_and_counts_rows(tmp_path: Path, extension: str) -> None:
    source_file = tmp_path / f"official{extension}"
    if extension == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(BEIJING_HOSPITALS.expected_headers))
        sheet.append(["医院甲"])
        sheet.append(["医院乙"])
        workbook.save(source_file)
    else:
        workbook_xls = xlwt.Workbook()
        sheet_xls = workbook_xls.add_sheet("Sheet1")
        sheet_xls.write(0, 0, "机构名称")
        sheet_xls.write(1, 0, "医院甲")
        sheet_xls.write(2, 0, "医院乙")
        workbook_xls.save(str(source_file))

    inspection = inspect_open_data_file(source_file, BEIJING_HOSPITALS)

    assert inspection.detected_format == extension.removeprefix(".")
    assert inspection.row_count == 2
    assert inspection.schema_match is True


def test_preflight_rejects_file_extension_content_mismatch(tmp_path: Path) -> None:
    source_file = tmp_path / "not-really.xlsx"
    source_file.write_text("机构名称\n医院甲\n", encoding="utf-8-sig")

    with pytest.raises(ValueError, match="format.*match"):
        inspect_open_data_file(source_file, BEIJING_HOSPITALS)


def test_beijing_official_ten_column_schema_is_recognized(tmp_path: Path) -> None:
    source_file = tmp_path / "beijing.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "北京市医疗保障局-定点医疗机构信息"
    sheet.append(BEIJING_HEADERS)
    sheet.append(
        (1, "测试医院", "11000001", "110101", "综合", "三级", "测试路", "x1", "created", "updated")
    )
    workbook.save(source_file)

    inspection = inspect_open_data_file(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)

    assert inspection.headers == BEIJING_HEADERS
    assert inspection.row_count == 1
    assert inspection.schema_match is True


def test_shenzhen_zip_preflight_reports_archive_and_member_fingerprints(tmp_path: Path) -> None:
    source_file = tmp_path / "baoan.zip"
    member_name = "宝安区-医院基本信息_2920002800636.xlsx"
    member_bytes = xlsx_bytes(
        "数据集1",
        SHENZHEN_HEADERS,
        [
            (
                "id1",
                "宝安区中心医院",
                "宝安区",
                "部门",
                "地址",
                "",
                "",
                "",
                "",
                "公立",
                "三级",
                "甲等",
                "是",
                "设有眼科门诊",
                "",
                "眼科",
                "公交",
            )
        ],
    )
    write_zip(source_file, {member_name: member_bytes})

    inspection = inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)

    assert inspection.detected_format == "zip"
    assert inspection.container_format == "zip"
    assert inspection.archive_member_name == member_name
    assert inspection.archive_member_sha256 == hashlib.sha256(member_bytes).hexdigest()
    assert inspection.archive_member_size_bytes == len(member_bytes)
    assert inspection.headers == SHENZHEN_HEADERS
    assert inspection.row_count == 1
    assert inspection.schema_match is True


@pytest.mark.parametrize(
    "member_name",
    ["../hospital.xlsx", "/hospital.xlsx", "folder/../../hospital.xlsx", "hospital.exe"],
)
def test_shenzhen_zip_rejects_unsafe_or_non_xlsx_member_paths(
    tmp_path: Path, member_name: str
) -> None:
    source_file = tmp_path / "baoan.zip"
    write_zip(source_file, {member_name: b"not a workbook"})

    with pytest.raises(ValueError):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_zip_rejects_encrypted_archive(tmp_path: Path) -> None:
    source_file = tmp_path / "encrypted.zip"
    member_name = "hospital.xlsx"
    member = xlsx_bytes("数据集1", SHENZHEN_HEADERS, [])
    write_zip(source_file, {member_name: member}, encrypted_flag=True)

    with pytest.raises(ValueError, match="encrypted"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_zip_rejects_multiple_data_members(tmp_path: Path) -> None:
    source_file = tmp_path / "multiple.zip"
    member = xlsx_bytes("数据集1", SHENZHEN_HEADERS, [])
    write_zip(source_file, {"one.xlsx": member, "two.xlsx": member})

    with pytest.raises(ValueError, match="one data file"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_zip_rejects_archive_without_an_xlsx_data_member(tmp_path: Path) -> None:
    source_file = tmp_path / "no-xlsx.zip"
    write_zip(source_file, {"readme.txt": b"not an XLSX"})

    with pytest.raises(ValueError, match="XLSX"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_zip_rejects_nested_zip_member(tmp_path: Path) -> None:
    source_file = tmp_path / "nested.zip"
    write_zip(source_file, {"nested.zip": b"PK\x03\x04"})

    with pytest.raises(ValueError, match="XLSX"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_zip_rejects_executable_member(tmp_path: Path) -> None:
    source_file = tmp_path / "executable.zip"
    with zipfile.ZipFile(source_file, "w") as archive:
        info = zipfile.ZipInfo("hospital.xlsx")
        info.external_attr = 0o100755 << 16
        archive.writestr(
            info,
            xlsx_bytes("数据集1", SHENZHEN_HEADERS, []),
        )

    with pytest.raises(ValueError, match="executable"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_zip_rejects_oversized_uncompressed_member(tmp_path: Path) -> None:
    source_file = tmp_path / "oversized.zip"
    write_zip(source_file, {"hospital.xlsx": b"x" * (10 * 1024 * 1024 + 1)})

    with pytest.raises(ValueError, match="10 MiB"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_shenzhen_preflight_fails_when_configured_sheet_is_missing(tmp_path: Path) -> None:
    source_file = tmp_path / "wrong-sheet.zip"
    member = xlsx_bytes("资源描述信息", SHENZHEN_HEADERS, [])
    write_zip(source_file, {"hospital.xlsx": member})

    with pytest.raises(ValueError, match="数据集1"):
        inspect_open_data_file(source_file, SHENZHEN_BAOAN_HOSPITALS)


def test_preflight_marks_malformed_beijing_administrative_code_not_ready(
    tmp_path: Path,
) -> None:
    source_file = tmp_path / "malformed-district.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "北京市医疗保障局-定点医疗机构信息"
    sheet.append(BEIJING_HEADERS)
    sheet.append(
        (1, "测试医院", "11000001", "东城区", "综合", "三级", "测试路", "x1", "created", "updated")
    )
    workbook.save(source_file)

    inspection = inspect_open_data_file(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)

    assert inspection.schema_match is False
    assert inspection.validation_errors == (
        "row 2: 所属区 must be a six-digit administrative code",
    )
