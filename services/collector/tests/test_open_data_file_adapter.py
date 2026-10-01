from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pytest
import xlwt
from openpyxl import Workbook

from eye_collector.etl.evidence import extract_ophthalmology_evidence
from eye_collector.etl.parser import parse_snapshot
from eye_collector.sources.open_data_file import (
    BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS,
    BEIJING_HOSPITALS,
    SHENZHEN_BAOAN_HOSPITALS,
    OpenDataFileAdapter,
)


def write_csv(path: Path, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)


def test_csv_adapter_maps_by_named_headers_and_preserves_source_values(tmp_path: Path) -> None:
    source_file = tmp_path / "hospitals.csv"
    write_csv(
        source_file,
        ("机构名称",),
        [("北京测试医院",), ("北京眼科医院",)],
    )
    adapter = OpenDataFileAdapter(source_file, BEIJING_HOSPITALS)

    records = list(adapter.iter_records("110000", limit=1))

    assert len(records) == 1
    assert records[0].source_key == "北京测试医院"
    assert records[0].raw_payload == {
        "name": "北京测试医院",
        "source_fields": {"机构名称": "北京测试医院"},
    }
    assert records[0].source_url == BEIJING_HOSPITALS.dataset_url
    assert adapter.descriptor.access_method == "file"


def test_adapter_fails_closed_when_any_expected_header_changes(tmp_path: Path) -> None:
    source_file = tmp_path / "designated.csv"
    write_csv(
        source_file,
        ("医院名称", "医院地址", "医院等级", "医院类别", "所属区", "定点医疗机构代码"),
        [("测试医院", "北京市东城区测试路", "三级", "综合", "东城区", "12345")],
    )
    adapter = OpenDataFileAdapter(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)

    with pytest.raises(ValueError, match="header schema"):
        list(adapter.iter_records("110000", limit=1))


def test_adapter_rejects_unexpected_header_and_duplicate_header(tmp_path: Path) -> None:
    for headers in (("机构名称", "电话"), ("机构名称", "机构名称")):
        source_file = tmp_path / "hospitals.csv"
        write_csv(source_file, headers, [("测试医院", "123")])
        adapter = OpenDataFileAdapter(source_file, BEIJING_HOSPITALS)
        with pytest.raises(ValueError, match="header schema"):
            list(adapter.iter_records("110000", limit=1))


def test_adapter_rejects_inconsistent_csv_row_width(tmp_path: Path) -> None:
    source_file = tmp_path / "designated.csv"
    write_csv(
        source_file,
        BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS.expected_headers,
        [("1", "测试医院", "11000001", "110101", "01", "03", "地址", "x1", "created")],
    )
    adapter = OpenDataFileAdapter(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)

    with pytest.raises(ValueError, match="row width"):
        list(adapter.iter_records("110000", limit=1))


def test_adapter_rejects_unsupported_file_type(tmp_path: Path) -> None:
    source_file = tmp_path / "hospitals.json"
    source_file.write_text("{}", encoding="utf-8")

    with pytest.raises(ValueError, match="supported formats"):
        OpenDataFileAdapter(source_file, BEIJING_HOSPITALS)


def test_adapter_caps_one_run_and_rejects_oversized_file(tmp_path: Path) -> None:
    source_file = tmp_path / "hospitals.csv"
    write_csv(source_file, ("机构名称",), [("测试医院",)])
    adapter = OpenDataFileAdapter(source_file, BEIJING_HOSPITALS)
    with pytest.raises(ValueError, match="explicit record limit"):
        list(adapter.iter_records("110000"))
    with pytest.raises(ValueError, match="150"):
        list(adapter.iter_records("110000", limit=151))

    source_file.write_bytes(b"x" * (10 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="file size"):
        OpenDataFileAdapter(source_file, BEIJING_HOSPITALS)


def test_xlsx_adapter_uses_exact_header_names_regardless_of_column_order(
    tmp_path: Path,
) -> None:
    source_file = tmp_path / "designated.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "北京市医疗保障局-定点医疗机构信息"
    sheet.append(list(reversed(BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS.expected_headers)))
    row = {
        "序号": 1,
        "医院名称": "北京市测试医院",
        "定点医疗机构编码": "01020304",
        "所属区": "110101",
        "医院类别": "01",
        "医院等级": "03",
        "医院地址": "北京市测试区测试路1号",
        "数据唯一记录号": "ignored-id",
        "数据创建时间": "2026-01-01",
        "数据更新时间": "2026-02-01",
    }
    sheet.append(
        [
            row[header]
            for header in reversed(BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS.expected_headers)
        ]
    )
    workbook.save(source_file)

    adapter = OpenDataFileAdapter(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)
    record = next(adapter.iter_records("110000", limit=1))

    assert record.source_key == "01020304"
    assert record.raw_payload["name"] == "北京市测试医院"
    assert record.raw_payload["address"] == "北京市测试区测试路1号"
    assert record.raw_payload["administrative_code"] == "110101"
    assert record.raw_payload["registration_id"] == "01020304"
    assert record.raw_payload["source_fields"]["医院等级"] == "03"
    assert set(record.raw_payload["source_fields"]) == {
        "医院名称",
        "医院地址",
        "所属区",
        "医院等级",
        "医院类别",
        "定点医疗机构编码",
    }
    assert "序号" not in record.raw_payload["source_fields"]
    assert "数据唯一记录号" not in record.raw_payload["source_fields"]


def test_xls_adapter_reads_text_identifiers_without_numeric_coercion(tmp_path: Path) -> None:
    source_file = tmp_path / "designated.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("北京市医疗保障局-定点医疗机构信息")
    headers = BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS.expected_headers
    row = {
        "序号": 1,
        "医院名称": "北京市测试医院",
        "医院地址": "北京市测试区测试路1号",
        "医院等级": "03",
        "医院类别": "01",
        "所属区": "110101",
        "定点医疗机构编码": "01020304",
        "数据唯一记录号": "ignored-id",
        "数据创建时间": "2026-01-01",
        "数据更新时间": "2026-02-01",
    }
    for column, header in enumerate(headers):
        sheet.write(0, column, header)
        sheet.write(1, column, row[header])
    workbook.save(str(source_file))

    adapter = OpenDataFileAdapter(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)
    record = next(adapter.iter_records("110000", limit=1))

    assert record.source_key == "01020304"
    assert record.raw_payload["registration_id"] == "01020304"
    assert record.raw_payload["administrative_code"] == "110101"


def test_shenzhen_strengths_field_can_supply_explicit_eye_evidence_without_contact_data(
    tmp_path: Path,
) -> None:
    source_file = tmp_path / "baoan.csv"
    row = {header: "" for header in SHENZHEN_BAOAN_HOSPITALS.expected_headers}
    row.update(
        {
            "文档ID": "row-1",
            "名称": "宝安测试医院",
            "所在区县": "宝安区",
            "详细地址": "深圳市宝安区测试路1号",
            "医疗优势与特长": "眼底照相设备",
            "医院简介": "设有眼科门诊",
            "联系电话": "075500000000",
            "电子邮箱": "private@example.test",
        }
    )
    write_csv(
        source_file,
        SHENZHEN_BAOAN_HOSPITALS.expected_headers,
        [tuple(row[header] for header in SHENZHEN_BAOAN_HOSPITALS.expected_headers)],
    )
    adapter = OpenDataFileAdapter(source_file, SHENZHEN_BAOAN_HOSPITALS)
    source_record = next(adapter.iter_records("440306", limit=1))
    parsed = parse_snapshot("source-record-1", source_record.raw_payload)

    assert parsed.record is not None
    extraction = extract_ophthalmology_evidence(parsed.record)
    assert extraction.status == "evidence_found"
    assert extraction.evidence[0].field_name == "hospital_description"
    assert extraction.evidence[0].evidence_text == "设有眼科门诊"
    assert "specialties" in source_record.raw_payload
    assert "hospital_description" in source_record.raw_payload
    assert "phone" not in source_record.raw_payload
    assert "email" not in source_record.raw_payload


def test_adapter_rejects_file_changed_after_fingerprint_was_captured(tmp_path: Path) -> None:
    source_file = tmp_path / "hospitals.csv"
    write_csv(source_file, ("机构名称",), [("原始医院",)])
    adapter = OpenDataFileAdapter(source_file, BEIJING_HOSPITALS)
    write_csv(source_file, ("机构名称",), [("替换后的医院",)])

    with pytest.raises(ValueError, match="changed after fingerprint"):
        list(adapter.iter_records("110000", limit=1))


def test_shenzhen_adapter_reads_configured_sheet_from_official_zip(tmp_path: Path) -> None:
    source_file = tmp_path / "baoan.zip"
    workbook_bytes = io.BytesIO()
    workbook = Workbook()
    workbook.active.title = "资源描述信息"
    workbook.active.append(("不是数据表",))
    sheet = workbook.create_sheet("数据集1")
    sheet.append(list(SHENZHEN_BAOAN_HOSPITALS.expected_headers))
    values = {header: "" for header in SHENZHEN_BAOAN_HOSPITALS.expected_headers}
    values.update(
        {
            "文档ID": "doc-001",
            "名称": "深圳市宝安区中心医院",
            "所在区县": "宝安区",
            "详细地址": "深圳市宝安区测试路",
            "医院简介": "医院设有眼科门诊。",
            "联系电话": "13800000000",
        }
    )
    sheet.append([values[header] for header in SHENZHEN_BAOAN_HOSPITALS.expected_headers])
    workbook.save(workbook_bytes)
    with zipfile.ZipFile(source_file, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("宝安区-医院基本信息_2920002800636.xlsx", workbook_bytes.getvalue())

    adapter = OpenDataFileAdapter(source_file, SHENZHEN_BAOAN_HOSPITALS)
    record = next(adapter.iter_records("440306", limit=1))

    assert adapter.original_filename == "baoan.zip"
    assert adapter.file_sha256 is not None
    assert adapter.archive_member_name == "宝安区-医院基本信息_2920002800636.xlsx"
    assert record.source_key == "doc-001"
    assert record.raw_payload["hospital_description"] == "医院设有眼科门诊。"
    assert record.raw_payload["source_fields"]["名称"] == "深圳市宝安区中心医院"
    assert "联系电话" not in record.raw_payload["source_fields"]


def test_beijing_adapter_rejects_malformed_nonempty_district_code(tmp_path: Path) -> None:
    source_file = tmp_path / "district.csv"
    write_csv(
        source_file,
        BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS.expected_headers,
        [("1", "测试医院", "reg-1", "东城区", "01", "03", "测试路", "x1", "created", "updated")],
    )

    adapter = OpenDataFileAdapter(source_file, BEIJING_DESIGNATED_MEDICAL_INSTITUTIONS)

    with pytest.raises(ValueError, match="six-digit administrative code"):
        list(adapter.iter_records("110000", limit=1))


def test_same_shenzhen_document_id_with_changed_content_remains_two_snapshots(
    tmp_path: Path,
) -> None:
    source_file = tmp_path / "baoan.csv"
    headers = SHENZHEN_BAOAN_HOSPITALS.expected_headers
    first = {header: "" for header in headers}
    first.update({"文档ID": "doc-1", "名称": "同名医院", "医院简介": "普通综合医院"})
    second = dict(first, **{"医院简介": "医院设有眼科门诊"})
    write_csv(
        source_file,
        headers,
        [
            tuple(first[header] for header in headers),
            tuple(second[header] for header in headers),
        ],
    )

    records = list(OpenDataFileAdapter(source_file, SHENZHEN_BAOAN_HOSPITALS).iter_records(
        "440306", limit=2
    ))

    assert [record.source_key for record in records] == ["doc-1", "doc-1"]
    assert records[0].raw_payload != records[1].raw_payload
    assert records[0].raw_payload["name"] == records[1].raw_payload["name"]
